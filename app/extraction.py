"""Extraction stages that touch the filesystem or subprocesses.

Each stage returns (data, error_message) instead of raising, so the endpoint
can degrade per stage: a file ExifTool chokes on still gets basics, detection,
and whatever supplements succeed.
"""

from __future__ import annotations

import hashlib
import json
import re
import shutil
import subprocess
import zipfile
from pathlib import Path

import puremagic

# -j -G1: JSON grouped by specific tag family. -a: keep duplicate tags.
# -u: include unknown tags. -ee: extract embedded streams (video, docs).
# -c "%+.6f": signed decimal GPS so the frontend can link to a map without
# parsing degrees/minutes/seconds.
EXIFTOOL_ARGS = ["-j", "-G1", "-a", "-u", "-ee", "-c", "%+.6f"]

# Generous ceiling: -ee on long videos is slow, but a hung perl process must
# not pin the request forever.
EXIFTOOL_TIMEOUT_S = 120


def exiftool_path() -> str | None:
    return shutil.which("exiftool")


def run_exiftool(path: Path) -> tuple[dict | None, str | None]:
    cmd = [exiftool_path(), *EXIFTOOL_ARGS, "--", str(path)]
    try:
        proc = subprocess.run(cmd, capture_output=True, text=True, timeout=EXIFTOOL_TIMEOUT_S)
    except subprocess.TimeoutExpired:
        return None, f"exiftool timed out after {EXIFTOOL_TIMEOUT_S}s"
    # ExifTool exits nonzero for recoverable problems while still emitting
    # valid JSON, so trust stdout first and only surface stderr as the error.
    if proc.stdout.strip():
        try:
            records = json.loads(proc.stdout)
            if records:
                return records[0], (proc.stderr.strip() or None)
        except json.JSONDecodeError:
            pass
    return None, proc.stderr.strip() or f"exiftool exited with code {proc.returncode} and no output"


def compute_hashes(path: Path) -> dict[str, str]:
    md5 = hashlib.md5()
    sha256 = hashlib.sha256()
    with open(path, "rb") as fh:
        while chunk := fh.read(1 << 20):
            md5.update(chunk)
            sha256.update(chunk)
    return {"md5": md5.hexdigest(), "sha256": sha256.hexdigest()}


def detect_type(path: Path) -> tuple[str | None, str | None, list[str]]:
    """Content-based detection. Returns (best extension, best mime, all
    candidate extensions). puremagic raises on no match; that is simply
    "unknown", not an error worth reporting.
    """
    try:
        matches = puremagic.magic_file(str(path))
    except Exception:
        return None, None, []
    if not matches:
        return None, None, []
    best = matches[0]
    ext = (best.extension or "").lstrip(".") or None
    candidates = [m.extension.lstrip(".") for m in matches if m.extension]
    return ext, best.mime_type or None, candidates


def _pdf_supplement(path: Path) -> dict:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    tags: dict = {"Encrypted": bool(reader.is_encrypted)}
    if reader.is_encrypted:
        # An empty user password (owner-locked PDF) is common; page details
        # are still worth extracting when it works.
        try:
            if not reader.decrypt(""):
                return tags
        except Exception:
            return tags
    tags["PageCount"] = len(reader.pages)
    sizes = []
    # Cap the per-page walk: past a few dozen pages the size list stops adding
    # information and huge PDFs get slow.
    for page in reader.pages[:50]:
        box = page.mediabox
        sizes.append(f"{float(box.width):g} x {float(box.height):g} pt")
    unique = sorted(set(sizes))
    if len(unique) == 1:
        tags["PageSize"] = unique[0]
    elif unique:
        tags["PageSizes"] = unique
    return tags


# ExifTool reads docProps/app.xml too, but only partially for some producers;
# these fields are cheap to re-extract and merge_supplements drops duplicates.
_APP_XML_FIELDS = ("Company", "Manager", "Application", "AppVersion", "TotalTime")


def _ooxml_app_fields(path: Path) -> dict:
    tags = {}
    with zipfile.ZipFile(path) as zf:
        if "docProps/app.xml" not in zf.namelist():
            return tags
        xml = zf.read("docProps/app.xml").decode("utf-8", "replace")
        for field in _APP_XML_FIELDS:
            m = re.search(rf"<{field}>([^<]+)</{field}>", xml)
            if m and m.group(1).strip():
                tags[field] = m.group(1).strip()
    return tags


def _docx_supplement(path: Path) -> dict:
    import docx

    document = docx.Document(str(path))
    props = document.core_properties
    tags = {
        "Creator": props.author or None,
        "LastModifiedBy": props.last_modified_by or None,
        "RevisionNumber": props.revision if props.revision else None,
        "CreateDate": props.created,
        "ModifyDate": props.modified,
    }
    tags.update(_ooxml_app_fields(path))
    with zipfile.ZipFile(path) as zf:
        names = zf.namelist()
        if "word/settings.xml" in names:
            settings = zf.read("word/settings.xml").decode("utf-8", "replace")
            # Track-changes recording currently on. ExifTool does not surface this.
            if "<w:trackChanges" in settings:
                tags["TrackChanges"] = True
        if "word/document.xml" in names:
            body = zf.read("word/document.xml").decode("utf-8", "replace")
            # Unaccepted insertions/deletions persist even after tracking is
            # switched off, so check the body independently of settings.
            if "<w:ins " in body or "<w:del " in body:
                tags["TrackedRevisionsInBody"] = True
    return {k: v for k, v in tags.items() if v is not None}


def _xlsx_supplement(path: Path) -> dict:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True)
    try:
        props = workbook.properties
        tags = {
            "Creator": props.creator or None,
            "LastModifiedBy": props.lastModifiedBy or None,
            "RevisionNumber": props.revision if props.revision else None,
            "CreateDate": props.created,
            "ModifyDate": props.modified,
        }
    finally:
        workbook.close()
    tags.update(_ooxml_app_fields(path))
    return {k: v for k, v in tags.items() if v is not None}


# Supplement group names double as the provenance marker in the merged output.
_SUPPLEMENTS = {
    "pdf": ("pypdf", _pdf_supplement),
    "docx": ("python-docx", _docx_supplement),
    "xlsx": ("openpyxl", _xlsx_supplement),
}


def run_supplements(path: Path, claimed_ext: str | None,
                    detected_ext: str | None) -> tuple[dict[str, dict], list[str]]:
    """Run format-specific extractors. Keyed on both the detected and claimed
    extension so a renamed .docx still gets the Office pass, and datetime
    values are ISO-stringified here so the merged dict stays JSON-safe.
    """
    supplements: dict[str, dict] = {}
    errors: list[str] = []
    wanted = {e for e in (claimed_ext, detected_ext) if e}
    for ext in wanted:
        entry = _SUPPLEMENTS.get(ext.lower())
        if not entry:
            continue
        group, func = entry
        if group in supplements:
            continue
        try:
            tags = func(path)
        except Exception as exc:
            errors.append(f"{group}: {type(exc).__name__}: {exc}")
            continue
        for key, value in tags.items():
            if hasattr(value, "isoformat"):
                tags[key] = value.isoformat(sep=" ")
        if tags:
            supplements[group] = tags
    return supplements, errors
